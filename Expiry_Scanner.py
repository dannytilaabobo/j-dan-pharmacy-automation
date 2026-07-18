print("=== J-DAN SYSTEM IS STARTING... ===")

import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
import pandas as pd
import datetime
import os
import shutil
import openpyxl
import re

# Anchor to this script's own folder, not the launch-time working directory
# (which can end up somewhere unexpected -- e.g. C:\Windows\System32 -- if
# launched via a shortcut, elevation, or scheduler that doesn't set cwd).
os.chdir(os.path.dirname(os.path.abspath(__file__)))

# ==========================================
#    AUDIT TRAIL / LOGBOOK SYSTEM
# ==========================================
print("\n" + "="*40)
print("     J-DAN EXPIRY REPORT SYSTEM")
print("="*40)
staff_name = input("Who is processing the pull-out? (Enter Staff Name): ")

current_time = datetime.datetime.now().strftime('%Y-%m-%d %I:%M %p')
with open("JDan_System_Logs.txt", "a") as logbook:
    logbook.write(f"[{current_time}] - USER: {staff_name.upper()} - Opened the EXPIRIES REPORT\n")

print(f"\nThank you, {staff_name.upper()}! Loading the application...\n")
# ==========================================

master_file = 'Inventory_Master.xlsx' 

main_folder = 'JDan_PullOut_System'
reports_folder = os.path.join(main_folder, '1_Excel_Reports')
proofs_folder = os.path.join(main_folder, '2_Archived_Proofs')
upload_box_folder = 'ATTACHED PROOF FOR EXPIRY PULL OUT'

for folder in [main_folder, reports_folder, proofs_folder, upload_box_folder]:
    if not os.path.exists(folder):
        os.makedirs(folder)

def standardize_date(date_str):
    if pd.isna(date_str) or str(date_str).strip() == "":
        return pd.NaT
    date_str = str(date_str).strip()
    if date_str.upper() in ["NO-EXP", "NAN", "NAT", "NONE"]:
        return pd.NaT

    match_mmyy = re.match(r'^(\d{1,2})[/-](\d{2})$', date_str)
    if match_mmyy:
        month = int(match_mmyy.group(1))
        year = int(match_mmyy.group(2)) + 2000 
        return pd.to_datetime(f"{year}-{month}-01") + pd.offsets.MonthEnd(0)

    match_mmyyyy = re.match(r'^(\d{1,2})[/-](\d{4})$', date_str)
    if match_mmyyyy:
        month = int(match_mmyyyy.group(1))
        year = int(match_mmyyyy.group(2))
        return pd.to_datetime(f"{year}-{month}-01") + pd.offsets.MonthEnd(0)
    
    try:
        return pd.to_datetime(date_str)
    except:
        return pd.NaT

def scan_expiries():
    try:
        df = pd.read_excel(master_file, sheet_name='INVENTORY MASTER')
        df.columns = df.columns.str.strip() 
        df['Current_Stock'] = pd.to_numeric(df['Current_Stock'], errors='coerce').fillna(0)
        df['Expiry_Date'] = df['Expiry_Date'].astype(str).str.strip()
        
        df['Expiry_Date_dt'] = df['Expiry_Date'].apply(standardize_date)
        
        today = pd.Timestamp.now().normalize()
        limit_date = today + pd.Timedelta(days=90) 
        
        ignore_blanks = ignore_blanks_var.get()

        def is_expiring_or_blank(row):
            if row['Current_Stock'] <= 0: return False
            val = str(row['Expiry_Date']).strip().upper()
            is_blank = val in ["", "NAN", "NAT", "NONE"]
            
            if is_blank: return not ignore_blanks
            if val == "NO-EXP": return False 
            if pd.isna(row['Expiry_Date_dt']): return not ignore_blanks
            if row['Expiry_Date_dt'] <= limit_date: return True
            return False

        mask = df.apply(is_expiring_or_blank, axis=1)
        expiring_df = df[mask].copy()
        
        expiring_df = expiring_df.drop_duplicates(subset=['Item_Code'])
        
        for item in tree.get_children(): tree.delete(item)
            
        if expiring_df.empty:
            messagebox.showinfo("System Status", "All clear. No medications are set to expire within the next 90 days.")
        else:
            for _, row in expiring_df.iterrows():
                tree.insert("", tk.END, values=(
                    "[ ]", str(row.get('Item_Code', '')), str(row['Product_Name']), 
                    str(row['Lot_Number']), "", "", "", "EXPIRY ALERT"
                ))
            messagebox.showinfo("Scan Complete", f"Found {len(expiring_df)} products that require physical shelf verification.")
    except Exception as e:
        messagebox.showerror("System Error", f"Error reading the Masterlist. Please check if the file is open or if the sheet name matches.\n\nDetails:\n{e}")

# MGA BAGONG SELECT ALL AT UNSELECT ALL FUNCTIONS
def select_all():
    for item in tree.get_children():
        vals = list(tree.item(item, 'values'))
        vals[0] = '[ X ]'
        tree.item(item, values=vals)

def unselect_all():
    for item in tree.get_children():
        vals = list(tree.item(item, 'values'))
        vals[0] = '[ ]'
        tree.item(item, values=vals)

def toggle_check(event):
    region = tree.identify("region", event.x, event.y)
    if region == "cell":
        col = tree.identify_column(event.x)
        if col == '#1': 
            item = tree.identify_row(event.y)
            if item:
                vals = list(tree.item(item, 'values'))
                vals[0] = '[ X ]' if vals[0] == '[ ]' else '[ ]'
                tree.item(item, values=vals)

def set_pullout_qty(event=None):
    selected_item = tree.selection()
    if not selected_item: return
    
    item = selected_item[0]
    vals = list(tree.item(item, 'values'))
    
    col = tree.identify_column(event.x)
    if col == '#1': return 
    
    if vals[7] == "MANUAL ADDITION":
        messagebox.showinfo("Info", "This is a manually added item. If you need to change its details, please process it or rescan.")
        return

    product = vals[2]
    
    actual_qty = simpledialog.askinteger("STEP 1/2: PULL OUT QUANTITY", 
        f"Item: {product}\n\nHow many units are you actually pulling out from the shelf?", 
        minvalue=0)
    
    if actual_qty is not None:
        vals[4] = str(actual_qty)
        vals[5] = "SYSTEM ALERT"

        next_date = simpledialog.askstring("STEP 2/2: TENTATIVE NEXT EXPIRY", 
            f"Item: {product}\n\nWhat is the TENTATIVE NEXT Expiry Date for the remaining items on the shelf?\n\n(Format: MM/YYYY, NO-EXP, or ALL EXPIRED)")
        
        if not next_date: next_date = "CHECK EXPIRY"
        vals[6] = str(next_date).upper()
        
        vals[0] = '[ X ]' 
        tree.item(item, values=vals)

def add_extra_item():
    search_window = tk.Toplevel(root)
    search_window.title("Add Extra Pull-Out Item")
    search_window.geometry("600x450")
    
    tk.Label(search_window, text="Search Product Name or Item Code:", font=("Helvetica", 10, "bold")).pack(pady=5)
    search_var = tk.StringVar()
    search_entry = tk.Entry(search_window, textvariable=search_var, width=50)
    search_entry.pack(pady=5)
    
    list_frame = tk.Frame(search_window)
    list_frame.pack(pady=5, fill="both", expand=True, padx=20)
    
    list_scroll = ttk.Scrollbar(list_frame, orient="vertical")
    list_scroll.pack(side="right", fill="y")
    
    result_list = tk.Listbox(list_frame, width=80, height=12, yscrollcommand=list_scroll.set, font=("Helvetica", 10))
    result_list.pack(side="left", fill="both", expand=True)
    list_scroll.config(command=result_list.yview)

    try:
        df_master = pd.read_excel(master_file, sheet_name='INVENTORY MASTER')
    except:
        messagebox.showerror("System Error", "Cannot read Inventory_Master.xlsx")
        search_window.destroy()
        return

    search_results_data = []

    def do_search():
        keyword = search_var.get().lower()
        result_list.delete(0, tk.END)
        search_results_data.clear()
        
        if not keyword: return
        
        matches = df_master[(df_master['Product_Name'].str.lower().str.contains(keyword, na=False) | 
                             df_master['Item_Code'].astype(str).str.lower().str.contains(keyword, na=False)) & 
                            (df_master['Current_Stock'] > 0)]
        
        unique_matches = matches.drop_duplicates(subset=['Item_Code'])
        
        for _, row in unique_matches.iterrows():
            display_text = f"{row['Item_Code']} | {row['Product_Name']}"
            result_list.insert(tk.END, display_text)
            search_results_data.append(row)

    def select_item():
        selection = result_list.curselection()
        if not selection: 
            messagebox.showwarning("Select Item", "Please select an item from the list.")
            return
        
        selected_idx = selection[0]
        row_data = search_results_data[selected_idx]
        product_name = row_data['Product_Name']
        
        actual_qty = simpledialog.askinteger("STEP 1/3: PULL OUT QUANTITY", 
            f"Item: {product_name}\n\nHow many units are you actually pulling out?", minvalue=1)
        
        if actual_qty is not None:
            pulled_date = simpledialog.askstring("STEP 2/3: PULLED ITEMS EXPIRY", 
                f"Item: {product_name}\n\nWhat is the Expiry Date printed on the {actual_qty} items you just pulled out?\n(Format: MM/YYYY or NO-EXP)")
            
            if not pulled_date: pulled_date = "CHECK EXPIRY"

            next_date = simpledialog.askstring("STEP 3/3: TENTATIVE NEXT EXPIRY", 
                f"Item: {product_name}\n\nWhat is the TENTATIVE NEXT Expiry Date for the remaining items on the shelf?\n(Format: MM/YYYY, NO-EXP, or ALL EXPIRED)")
            
            if not next_date: next_date = "CHECK EXPIRY"

            tree.insert("", tk.END, values=(
                "[ X ]", str(row_data['Item_Code']), str(product_name), "MIXED", 
                str(actual_qty), str(pulled_date).upper(), str(next_date).upper(), "MANUAL ADDITION"
            ))
            messagebox.showinfo("Success", "Extra item added to the main list and automatically ticked.")
            search_window.destroy()

    tk.Button(search_window, text="Search", command=do_search, width=20).pack(pady=2)
    tk.Button(search_window, text="Select & Add Item", command=select_item, font=("Helvetica", 10, "bold"), bg="#5bc0de", fg="white", width=25).pack(pady=10)

def generate_report():
    data = []
    for item in tree.get_children():
        vals = tree.item(item, 'values')
        if vals[0] == '[ X ]':
            data.append(list(vals[1:8]) + [staff_name.upper()]) 
    
    if not data:
        messagebox.showwarning("No Items Selected", "There are no ticked items.\nPlease tick [ X ] the items you want to include in the report before generating.")
        return
        
    df_report = pd.DataFrame(data, columns=['Item_Code', 'Product_Name', 'Lot_Number', 'PullOut_Qty', 'Pulled_Expiry', 'Next_Expiry', 'Status', 'Generated_By'])
    report_path = os.path.join(reports_folder, f'EXPIRY_CHECKLIST_{staff_name.upper()}_{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx')
    df_report.to_excel(report_path, index=False)
    messagebox.showinfo("SUCCESS", "Filtered Checklist Generated!\nThis report strictly contains ONLY the items marked with an [ X ].")

def process_pullout():
    items_to_process = []
    for item in tree.get_children():
        vals = tree.item(item, 'values')
        if vals[0] == '[ X ]':
            if vals[4] == "": 
                messagebox.showwarning("Incomplete Action", f"You checked the item '{vals[2]}' but did not update its quantities.\n\nPlease double-click it to input the data.")
                return
            items_to_process.append({
                'code': vals[1],
                'name': vals[2],
                'actual_qty': float(vals[4]),
                'pulled_expiry': vals[5],
                'next_expiry': vals[6]
            })

    if not items_to_process:
        messagebox.showwarning("No Items Selected", "No items were selected for processing.\nPlease tick the [ X ] box for the items you want to pull out.")
        return
    
    uploaded_files = [f for f in os.listdir(upload_box_folder) if f.lower().endswith(('.jpg', '.png', '.jpeg', '.pdf'))]
    if not uploaded_files:
        messagebox.showerror("Missing Evidence", "No proof attached! Please upload a photo of the expired medications to the designated folder before processing.")
        return

    try:
        wb = openpyxl.load_workbook(master_file)
        ws = wb['INVENTORY MASTER']
        
        headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
        
        try:
            item_code_col = headers.index('Item_Code') + 1
            stock_col = headers.index('Current_Stock') + 1
            expiry_col = headers.index('Expiry_Date') + 1
        except ValueError as e:
            messagebox.showerror("Header Format Error", f"Cannot find exact columns in the Excel file. Please ensure 'Item_Code', 'Current_Stock', and 'Expiry_Date' exist.\n\nError: {e}")
            return

        def parse_date_for_openpyxl(d_val):
            return standardize_date(d_val)

        log_entries = []
        final_report_data = [] 

        for item in items_to_process:
            qty_to_deduct = item['actual_qty']
            next_expiry = item['next_expiry']
            item_rows = []
            
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(row=r, column=item_code_col).value).strip() == str(item['code']).strip():
                    stock = ws.cell(row=r, column=stock_col).value
                    if stock and float(stock) > 0:
                        parsed_dt = parse_date_for_openpyxl(ws.cell(row=r, column=expiry_col).value)
                        if pd.isna(parsed_dt):
                            parsed_dt = datetime.datetime(2099, 12, 31)
                            
                        item_rows.append({
                            'row': r, 
                            'expiry_dt': parsed_dt, 
                            'stock': float(stock)
                        })
            
            item_rows.sort(key=lambda x: x['expiry_dt'])
            
            for ir in item_rows:
                available = ir['stock']
                
                if qty_to_deduct > 0:
                    if available >= qty_to_deduct:
                        new_stock = available - qty_to_deduct
                        ws.cell(row=ir['row'], column=stock_col).value = new_stock
                        qty_to_deduct = 0
                        
                        if new_stock > 0 and next_expiry not in ["", "ALL EXPIRED", "CHECK EXPIRY"]:
                            ws.cell(row=ir['row'], column=expiry_col).value = next_expiry
                    else:
                        ws.cell(row=ir['row'], column=stock_col).value = 0
                        qty_to_deduct -= available
                else:
                    if next_expiry not in ["", "ALL EXPIRED", "CHECK EXPIRY"]:
                        ws.cell(row=ir['row'], column=expiry_col).value = next_expiry
            
            log_entries.append(f"   -> {item['code']} | {item['name']} | Pulled Qty: {item['actual_qty']} | Pulled Date: {item['pulled_expiry']} | Tentative Next Expiry: {item['next_expiry']}")
            
            final_report_data.append([
                item['code'], item['name'], item['actual_qty'], item['pulled_expiry'], item['next_expiry'], staff_name.upper()
            ])

        wb.save(master_file)
        
        timestamp_str = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        
        timestamp = datetime.datetime.now().strftime('%Y-%m-%d %I:%M %p')
        with open("JDan_System_Logs.txt", "a") as logbook:
            logbook.write(f"[{timestamp}] - PULL-OUT PROCESS (FIFO ENGINE) by {staff_name.upper()}\n")
            for entry in log_entries:
                logbook.write(entry + "\n")

        df_completed = pd.DataFrame(final_report_data, columns=['Item_Code', 'Product_Name', 'Actual_Pulled_Qty', 'Pulled_Expiry_Date', 'Next_Tentative_Expiry', 'Processed_By'])
        completed_report_path = os.path.join(reports_folder, f'COMPLETED_PULLOUT_{staff_name.upper()}_{timestamp_str}.xlsx')
        df_completed.to_excel(completed_report_path, index=False)

        for f in uploaded_files:
            src = os.path.join(upload_box_folder, f)
            dst = os.path.join(proofs_folder, f"{timestamp_str}_{staff_name}_{f}")
            shutil.move(src, dst)

        messagebox.showinfo("Operation Successful", f"Masterlist Reconciled Successfully!\n\nA final report named 'COMPLETED_PULLOUT_{staff_name.upper()}_{timestamp_str}.xlsx' has been automatically generated and saved in the Reports folder.")
        scan_expiries() 
        
    except Exception as e:
        messagebox.showerror("System Error", f"Failed to update Masterlist. Please ensure Inventory_Master.xlsx is closed.\n\nDetails: {e}")

# --- GUI CONSTRUCTION ---
root = tk.Tk()
root.title("J-DAN EXPIRY REPORT SYSTEM")
root.geometry("1180x720")
root.configure(bg="#f4f4f4")

tk.Label(root, text="🚨 J-DAN EXPIRY REPORT SYSTEM 🚨", font=("Helvetica", 16, "bold"), fg="#d9534f", bg="#f4f4f4").pack(pady=10)

top_frame = tk.Frame(root, bg="#f4f4f4")
top_frame.pack(pady=5)
tk.Button(top_frame, text="1. SCAN ALERTS", font=("Helvetica", 10, "bold"), bg="#5bc0de", fg="white", width=18, command=scan_expiries).grid(row=0, column=0, padx=5)
tk.Button(top_frame, text="2. GENERATE REPORT", font=("Helvetica", 10, "bold"), bg="#f0ad4e", fg="white", width=20, command=generate_report).grid(row=0, column=1, padx=5)
tk.Button(top_frame, text="3. ADD EXTRA ITEM", font=("Helvetica", 10, "bold"), bg="#5cb85c", fg="white", width=18, command=add_extra_item).grid(row=0, column=2, padx=5)
tk.Button(top_frame, text="4. PROCESS PULL-OUT", font=("Helvetica", 10, "bold"), bg="#d9534f", fg="white", width=20, command=process_pullout).grid(row=0, column=3, padx=5)

mid_frame = tk.Frame(root, bg="#f4f4f4")
mid_frame.pack(pady=5, anchor="w", padx=20)

tk.Label(mid_frame, text="Instruction: Double-click an item to process. Pop-ups will automatically guide you through Qty and Next Expiry.", font=("Helvetica", 9, "italic"), bg="#f4f4f4").grid(row=0, column=0, padx=5)

ignore_blanks_var = tk.BooleanVar(value=True) 
tk.Checkbutton(mid_frame, text="Ignore Blank Expiry Dates", variable=ignore_blanks_var, bg="#f4f4f4", font=("Helvetica", 9, "bold")).grid(row=0, column=1, padx=20)

# MGA BAGONG SELECT ALL / UNSELECT ALL BUTTONS
btn_frame = tk.Frame(root, bg="#f4f4f4")
btn_frame.pack(pady=2, anchor="w", padx=20)
tk.Button(btn_frame, text="☑ Select All", font=("Helvetica", 9, "bold"), bg="#5cb85c", fg="white", command=select_all, width=15).grid(row=0, column=0, padx=5)
tk.Button(btn_frame, text="☐ Unselect All", font=("Helvetica", 9, "bold"), bg="#d9534f", fg="white", command=unselect_all, width=15).grid(row=0, column=1, padx=5)

tree_frame = tk.Frame(root, bg="#f4f4f4")
tree_frame.pack(pady=5, padx=20, fill="both", expand=True)

tree_scroll = ttk.Scrollbar(tree_frame, orient="vertical")
tree_scroll.pack(side="right", fill="y")

columns = ("Check", "Code", "Product", "Lot", "PullOut_Qty", "Pulled_Expiry", "Next_Expiry", "Status")
display_cols = ("Check", "Code", "Product", "Lot", "PullOut_Qty", "Pulled_Expiry", "Next_Expiry", "Status")

tree = ttk.Treeview(tree_frame, columns=columns, displaycolumns=display_cols, show="headings", height=15, yscrollcommand=tree_scroll.set)
tree_scroll.config(command=tree.yview)

for col in display_cols: tree.heading(col, text=col)

tree.column("Check", width=60, anchor='center')
tree.column("Code", width=90, anchor='center')
tree.column("Product", width=230)
tree.column("Lot", width=90, anchor='center')
tree.column("PullOut_Qty", width=90, anchor='center') 
tree.column("Pulled_Expiry", width=120, anchor='center') 
tree.column("Next_Expiry", width=120, anchor='center')
tree.column("Status", width=130, anchor='center')

tree.pack(side="left", fill="both", expand=True)

tree.bind('<Double-1>', set_pullout_qty)
tree.bind('<ButtonRelease-1>', toggle_check)

root.mainloop()
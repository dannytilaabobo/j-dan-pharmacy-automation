import openpyxl
import datetime
import tkinter as tk
from tkinter import filedialog

print("=== J-DAN BULK DATE FIXER ===")

# Buksan ang File Selector window
root = tk.Tk()
root.withdraw() # Itago ang main tkinter window para pop-up agad

print("Please select the Excel file in the window that appears...")
master_file = filedialog.askopenfilename(
    title="J-DAN: Select your Inventory Excel File",
    filetypes=[("Excel files", "*.xlsx")]
)

if not master_file:
    print("\nNo file was selected. Cancelling the process.")
else:
    print(f"\nFile selected: {master_file}")
    print("Fixing the dates that Excel altered...\n")
    try:
        # Siguraduhing nakasara ang file bago ito mag-run
        wb = openpyxl.load_workbook(master_file)

        # Gamitin ang tamang sheet name ng J-DAN Inventory
        ws = wb['INVENTORY MASTER']

        # Hanapin kung nasaan ang Expiry_Date column
        headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
        expiry_col = headers.index('Expiry_Date') + 1

        fixed_count = 0

        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=expiry_col)
            val = cell.value

            # Kapag ang value ay naging datetime object (auto-format ni Excel)
            if isinstance(val, datetime.datetime):
                # Kung ang Year ay 2024, 2025, o 2026, at ang Day ay 24 pataas:
                if val.year in [2024, 2025, 2026] and val.day >= 24:
                    correct_month = val.month
                    correct_year = 2000 + val.day # Yung araw (ex. 28) ay gagawing taon (2028)

                    # Ibabalik sa text format na "MM/YYYY" para hindi na magalaw
                    new_val = f"{correct_month:02d}/{correct_year}"
                    cell.value = new_val
                    fixed_count += 1

        wb.save(master_file)
        print(f"SUCCESS! 🎉 Automatically fixed {fixed_count} item(s) in your Excel file.")
        print("You may now open your Excel file to check the dates.")

    except Exception as e:
        print(f"ERROR: Could not process the file. Please make sure Excel is CLOSED before running this.\nDetails: {e}")

input("\nPress Enter to close this window...")
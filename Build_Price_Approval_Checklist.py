import glob
import datetime
import pandas as pd
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MASTER_FILE = 'Inventory_Master.xlsx'
BITPOS_GLOB = 'bitpos_sales_pricing_*.csv'

# Manual verification notes from the 2026-07-11 investigation. These do not
# get rediscovered automatically -- fill in new ones here if this checklist
# is regenerated for a different comparison run.
VERIFICATION_NOTES = {
    '66252450': ('CONFIRMED REAL PRICE DROP', 'POS already sold this item at P115 on 2026-06-28 -- '
                                                'not a typo, Master is the outdated one.'),
}
DELISTED_NOTES = {
    '37776367': ('GENUINE REPACK/DISCONTINUATION', 'Supplier replaced with SILVER SULFADIAZINE 10G '
                 '(code 26246000, already in Master at P160, 0 stock). 5 units of the 20G left on shelf -- sell through or return.'),
    '77925626': ('GENUINE REPACK/DISCONTINUATION', 'Supplier replaced with FUSIDIC ACID TOPICAL CRM 10G '
                 '(code 10320655, already in Master at P180, 0 stock). 5 units of the 5G left on shelf -- sell through or return.'),
}


def load_bitpos():
    path = sorted(glob.glob(BITPOS_GLOB))[-1]
    df = pd.read_csv(path, dtype=str)
    df.columns = [c.strip() for c in df.columns]
    df['ITEM CODE'] = df['ITEM CODE'].astype(str).str.strip()
    df['Retail'] = pd.to_numeric(df['Retail'], errors='coerce')
    return df


def load_master():
    df = pd.read_excel(MASTER_FILE, sheet_name='INVENTORY MASTER')
    df['Item_Code'] = df['Item_Code'].astype(str).str.strip()
    df['Retail_Price'] = pd.to_numeric(df['Retail_Price'], errors='coerce')
    return df


def style_header(ws, ncols):
    fill = PatternFill(start_color='154360', end_color='154360', fill_type='solid')
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    df_bitpos = load_bitpos()
    df_master = load_master()

    master_prices = df_master.groupby('Item_Code').agg(
        Product_Name=('Product_Name', 'first'),
        Master_Retail=('Retail_Price', 'first'),
        Current_Stock=('Current_Stock', 'sum'),
    ).reset_index()

    merged = df_bitpos.merge(master_prices, left_on='ITEM CODE', right_on='Item_Code', how='outer', indicator=True)

    both = merged[merged['_merge'] == 'both'].copy()
    both['price_diff'] = both['Retail'] - both['Master_Retail']
    mismatches = both[both['price_diff'].abs() > 0.001].copy()
    mismatches = mismatches.sort_values('price_diff', key=lambda s: s.abs(), ascending=False)

    delisted = merged[merged['_merge'] == 'right_only'][['Item_Code', 'Product_Name', 'Master_Retail', 'Current_Stock']].copy()

    out_path = f'JDan_Executive_Reports/PRICE_APPROVAL_CHECKLIST_{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        wb = writer.book

        # --- Sheet 1: Price update approvals ---
        rows = []
        for _, r in mismatches.iterrows():
            code = r['ITEM CODE']
            verif_status, verif_note = VERIFICATION_NOTES.get(code, ('NOT YET VERIFIED', ''))
            rows.append({
                'Item_Code': code,
                'Description': r['DESCRIPTION'],
                'Master Price (current)': r['Master_Retail'],
                'BITPOS Price (new)': r['Retail'],
                'Difference': r['price_diff'],
                'Verification Status': verif_status,
                'Verification Note': verif_note,
                'Decision (APPROVE / HOLD / SKIP)': '',
                'Reviewer Notes': '',
            })
        approval_df = pd.DataFrame(rows)
        approval_df.to_excel(writer, sheet_name='Price_Updates', index=False)
        ws1 = wb['Price_Updates']
        style_header(ws1, len(approval_df.columns))
        autosize(ws1, [12, 42, 14, 14, 11, 20, 45, 26, 30])
        dv1 = DataValidation(type='list', formula1='"APPROVE,HOLD,SKIP"', allow_blank=True)
        ws1.add_data_validation(dv1)
        dv1.add(f'H2:H{len(approval_df) + 1}')

        # --- Sheet 2: Delisted item verification ---
        rows2 = []
        for _, r in delisted.iterrows():
            code = r['Item_Code']
            verif_status, verif_note = DELISTED_NOTES.get(code, ('NOT YET VERIFIED', ''))
            rows2.append({
                'Item_Code': code,
                'Product_Name (Master)': r['Product_Name'],
                'Master Price': r['Master_Retail'],
                'Current Stock (Master)': r['Current_Stock'],
                'Verification Status': verif_status,
                'Verification Note': verif_note,
                'Decision (DISCONTINUE / KEEP / INVESTIGATE)': '',
                'Reviewer Notes': '',
            })
        delisted_df = pd.DataFrame(rows2)
        delisted_df.to_excel(writer, sheet_name='Delisted_Items', index=False)
        ws2 = wb['Delisted_Items']
        style_header(ws2, len(delisted_df.columns))
        autosize(ws2, [12, 32, 12, 14, 26, 55, 30, 30])
        dv2 = DataValidation(type='list', formula1='"DISCONTINUE,KEEP,INVESTIGATE"', allow_blank=True)
        ws2.add_data_validation(dv2)
        dv2.add(f'G2:G{len(delisted_df) + 1}')

    print(f'Checklist saved to: {out_path}')
    print(f'  Sheet "Price_Updates": {len(approval_df)} items awaiting APPROVE/HOLD/SKIP decision')
    print(f'  Sheet "Delisted_Items": {len(delisted_df)} items awaiting DISCONTINUE/KEEP/INVESTIGATE decision')


if __name__ == '__main__':
    main()

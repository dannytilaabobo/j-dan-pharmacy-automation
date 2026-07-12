import os
import shutil
import datetime
import pandas as pd
import openpyxl

from name_cleaning import clean_text, is_date_corrupted

MASTER_FILE = 'Inventory_Master.xlsx'
BACKUP_DIR = 'JDan_Master_Backups'
REPORT_DIR = 'JDan_Executive_Reports'


def backup_master():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    backup_path = os.path.join(BACKUP_DIR, f'Inventory_Master_BEFORE_NAME_CLEAN_{ts}.xlsx')
    shutil.copy2(MASTER_FILE, backup_path)
    return backup_path


def scan_date_corrupted_names():
    """Detect Product_Name cells Excel auto-converted to a date. Unlike
    Expiry_Date (Auto_Fix_Dates.py), there is no reliable rule to reconstruct
    arbitrary product-name text from a date value, so these are flagged for
    manual fix rather than guessed."""
    wb = openpyxl.load_workbook(MASTER_FILE)
    ws = wb['INVENTORY MASTER']
    headers = [str(c.value).strip() if c.value else '' for c in ws[1]]
    name_col = headers.index('Product_Name') + 1
    code_col = headers.index('Item_Code') + 1

    flagged = []
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=name_col).value
        if is_date_corrupted(val):
            flagged.append({
                'row': r,
                'Item_Code': ws.cell(row=r, column=code_col).value,
                'corrupted_value': str(val),
            })
    return flagged


def main():
    backup_path = backup_master()
    print(f'Backup created: {backup_path}')

    date_corrupted = scan_date_corrupted_names()
    if date_corrupted:
        print(f'\nWARNING: {len(date_corrupted)} Product_Name cells were auto-converted to dates by Excel.')
        print('These need MANUAL fixing (no safe auto-reconstruction rule exists for names):')
        for f in date_corrupted:
            print(f"  Row {f['row']} | Item_Code {f['Item_Code']} | corrupted value: {f['corrupted_value']}")
    else:
        print('\nNo date-corrupted Product_Name cells found.')

    df = pd.read_excel(MASTER_FILE, sheet_name='INVENTORY MASTER')

    changes = []
    cleaned_names = []
    for _, row in df.iterrows():
        old_name = row['Product_Name']
        if is_date_corrupted(old_name):
            cleaned_names.append(old_name)
            continue
        new_name = clean_text(old_name)
        cleaned_names.append(new_name)
        if new_name != old_name:
            changes.append({
                'Item_Code': row['Item_Code'],
                'Before': old_name,
                'After': new_name,
            })
    df['Product_Name'] = cleaned_names

    if changes:
        with pd.ExcelWriter(MASTER_FILE, engine='openpyxl', mode='w') as writer:
            df.to_excel(writer, sheet_name='INVENTORY MASTER', index=False)
        print(f'\n{len(changes)} Product_Name cell(s) cleaned and saved to {MASTER_FILE}.')

        os.makedirs(REPORT_DIR, exist_ok=True)
        report_path = os.path.join(
            REPORT_DIR, f'NAME_CLEANUP_REPORT_{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx')
        with pd.ExcelWriter(report_path, engine='openpyxl') as writer:
            pd.DataFrame(changes).to_excel(writer, sheet_name='Before_After', index=False)
            if date_corrupted:
                pd.DataFrame(date_corrupted).to_excel(writer, sheet_name='Needs_Manual_Fix', index=False)
        print(f'Before/after report saved to: {report_path}')
    else:
        print('\nNo Product_Name cells needed cleaning. Master file left untouched.')


if __name__ == '__main__':
    main()
